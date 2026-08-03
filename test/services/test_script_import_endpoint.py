import io
import os

import docx
from fastapi.testclient import TestClient

from app.asgi import app


def _build_package(tmp_path):
    document = docx.Document()
    document.add_heading("Sample Video", level=1)
    document.add_heading("4. Full Script", level=2)
    document.add_paragraph("Narration only. Scene markers are not read aloud.")
    document.add_paragraph("[SCENE 01]   00:00")
    document.add_paragraph("This is the first scene narration.")
    document.add_paragraph("[SCENE 02]   00:48")
    document.add_paragraph("This is the second scene narration.")

    document.add_heading("5. Scene Table", level=2)
    table = document.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Scene"
    table.rows[0].cells[1].text = "Timestamp"
    table.rows[0].cells[2].text = "Scene summary"
    table.rows[0].cells[3].text = "Image prompt (16:9)"

    row = table.add_row()
    row.cells[0].text = "[SCENE 01]"
    row.cells[1].text = "00:00"
    row.cells[2].text = "Opening scene"
    row.cells[3].text = "a wide establishing shot"

    row = table.add_row()
    row.cells[0].text = "[SCENE 02]"
    row.cells[1].text = "00:48"
    row.cells[2].text = "Second scene"
    row.cells[3].text = "a close up shot"

    path = os.path.join(tmp_path, "package.docx")
    document.save(path)
    return path


def test_import_script_docx_returns_script_and_scenes(tmp_path):
    docx_path = _build_package(tmp_path)

    with open(docx_path, "rb") as f:
        response = TestClient(app).post(
            "/api/v1/scripts/import",
            files={
                "file": (
                    "package.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == 200

    data = body["data"]
    assert "first scene narration" in data["video_script"]
    assert "second scene narration" in data["video_script"]
    assert len(data["scenes"]) == 2

    first_scene = data["scenes"][0]
    assert first_scene["scene_id"] == "01"
    assert first_scene["planned_start_seconds"] == 0
    assert first_scene["image_prompt"] == "a wide establishing shot"
    assert first_scene["summary"] == "Opening scene"


def test_import_script_docx_rejects_non_docx_extension():
    response = TestClient(app).post(
        "/api/v1/scripts/import",
        files={"file": ("script.txt", io.BytesIO(b"not a docx"), "text/plain")},
    )

    assert response.status_code == 400


def test_import_script_docx_rejects_docx_without_script_section(tmp_path):
    document = docx.Document()
    document.add_heading("No script section here", level=1)
    document.add_paragraph("Just some text.")
    path = os.path.join(tmp_path, "no_script.docx")
    document.save(path)

    with open(path, "rb") as f:
        response = TestClient(app).post(
            "/api/v1/scripts/import",
            files={
                "file": (
                    "no_script.docx",
                    f,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

    assert response.status_code == 400
