import os

import docx
import pytest

from app.services.script_import import (
    ScriptImportError,
    parse_docx_script,
    parse_timestamp,
)


def _build_package(tmp_path, include_table=True, include_markers=True):
    document = docx.Document()
    document.add_heading("Sample Video", level=1)
    document.add_heading("4. Full Script", level=2)
    document.add_paragraph("Narration only. Scene markers are not read aloud.")

    if include_markers:
        document.add_paragraph("[SCENE 01]   00:00")
    document.add_paragraph("This is the first scene narration.")
    document.add_paragraph("It has two sentences.")

    if include_markers:
        document.add_paragraph("[SCENE 02]   00:48")
    document.add_paragraph("This is the second scene narration.")

    document.add_heading("5. Scene Table", level=2)
    if include_table:
        table = document.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "Scene"
        table.rows[0].cells[1].text = "Timestamp"
        table.rows[0].cells[2].text = "Scene summary"
        table.rows[0].cells[3].text = "Image prompt (16:9)"

        row = table.add_row()
        row.cells[0].text = "[SCENE 01]"
        row.cells[1].text = "00:00"
        row.cells[2].text = "Opening scene"
        row.cells[3].text = "a wide establishing shot"

        row = table.add_row()
        row.cells[0].text = "[SCENE 02]"
        row.cells[1].text = "00:48"
        row.cells[2].text = "Second scene"
        row.cells[3].text = "a close up shot"

    path = os.path.join(tmp_path, "package.docx")
    document.save(path)
    return path


def test_parse_timestamp():
    assert parse_timestamp("00:48") == 48
    assert parse_timestamp("01:02:03") == 3723
    assert parse_timestamp("") is None
    assert parse_timestamp("not a timestamp") is None


def test_parse_docx_script_extracts_scenes_and_script(tmp_path):
    path = _build_package(tmp_path)
    result = parse_docx_script(path)

    assert len(result.scenes) == 2

    first, second = result.scenes
    assert first.scene_id == "01"
    assert first.planned_start_seconds == 0
    assert "first scene narration" in first.narration
    assert "two sentences" in first.narration
    assert first.image_prompt == "a wide establishing shot"
    assert first.summary == "Opening scene"

    assert second.scene_id == "02"
    assert second.planned_start_seconds == 48
    assert "second scene narration" in second.narration
    assert second.image_prompt == "a close up shot"

    assert "first scene narration" in result.video_script
    assert "second scene narration" in result.video_script
    # Editorial notes and scene markers must not leak into the narration.
    assert "Narration only" not in result.video_script
    assert "[SCENE" not in result.video_script


def test_parse_docx_script_without_table_still_returns_scenes(tmp_path):
    path = _build_package(tmp_path, include_table=False)
    result = parse_docx_script(path)

    assert len(result.scenes) == 2
    assert result.scenes[0].image_prompt == ""
    assert result.scenes[0].summary == ""


def test_parse_docx_script_without_script_heading_raises(tmp_path):
    document = docx.Document()
    document.add_heading("No script section here", level=1)
    document.add_paragraph("Just some text.")
    path = os.path.join(tmp_path, "no_heading.docx")
    document.save(path)

    with pytest.raises(ScriptImportError):
        parse_docx_script(path)


def test_parse_docx_script_without_scene_markers_raises(tmp_path):
    path = _build_package(tmp_path, include_markers=False)

    with pytest.raises(ScriptImportError):
        parse_docx_script(path)
